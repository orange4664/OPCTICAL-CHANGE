"""Build the Chinese main-thesis style DOCX deliverable."""

from __future__ import annotations

from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "output" / "doc"
PRIMARY_DOC = OUTPUT_DIR / "双栅WSe2复SHG零点与相位绕转_主论稿.docx"
ASCII_DOC = OUTPUT_DIR / "wse2_complex_shg_zero_phase_winding_main.docx"
FIG_DIR = ROOT / "output" / "figures"


def add_paragraph(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(21)
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.add_run(text)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_figure(doc: Document, filename: str, caption: str) -> None:
    path = FIG_DIR / filename
    if path.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(str(path), width=Inches(6.1))
    add_paragraph(doc, caption)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "SimSun"
    normal.font.size = Pt(10.5)
    for style_name in ["Heading 1", "Heading 2", "Title"]:
        style = doc.styles[style_name]
        style.font.name = "SimHei"


def build_document() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("双栅 WSe2 复数二次谐波场零点与相位绕转的数值认证")
    run.bold = True
    run.font.size = Pt(18)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(
        "Gate-programmable zeros and phase winding of the scalar complex SHG field in exciton-polaron-mediated bilayer WSe2"
    )
    run.italic = True
    run.font.size = Pt(11)

    doc.add_heading("摘要", level=1)
    add_paragraph(
        doc,
        "本文围绕 Cha 等人在双栅双层 WSe2 中通过 layer-dependent exciton-polaron effect 增强共振二次谐波产生的工作，进一步细化一个新的可验证创新点：双栅调控 exciton-polaron 共振，实现双层 WSe2 标量复 SHG 场的电控零点与相位绕转。本文不把普通 SHG 强度增强作为创新本身，而是把研究对象从强度调制推进到复数二次谐波场 E2w 的实部、虚部、相位和零点结构。",
    )
    add_paragraph(
        doc,
        "我们构建了一个 literature-constrained phenomenological Python 模型：先将顶栅和底栅电压映射到载流子密度 n 与位移场 D，再用层分辨复 Lorentz 振子描述 neutral exciton 与 exciton-polaron 通道，最后通过 2H 双层的相干场求和得到标量复 SHG 场。模型数值认证了在合理参数范围内，复 SHG 场可以被门压推近相消零点，并在围绕零点的 gate loop 上产生相位绕转。相位分辨 heterodyne SHG 可作为实验验证路径。",
    )

    doc.add_heading("English central claim", level=1)
    add_paragraph(
        doc,
        "We propose and numerically validate a gate-programmable scalar complex-SHG-field zero in exciton-polaron-mediated bilayer WSe2. The central advance is not stronger SHG intensity, but electrical control of the phase and topology of the emitted second-harmonic field.",
    )

    doc.add_heading("1. 主论文基线与创新边界", level=1)
    add_paragraph(
        doc,
        "主论文基线是 Cha et al. 2024 的 dual-gated bilayer WSe2 resonant SHG 工作[1]。该论文的关键事实是：双栅结构可以使空穴在双层中产生层选择性分布，使某一层进入 exciton-polaron 态，而另一层更接近 neutral-exciton 响应。由于上下层的共振条件和非线性响应不再等价，原本接近反演对称抵消的 2H 双层可以产生显著的 resonant SHG 增强或淬灭。",
    )
    add_paragraph(
        doc,
        "因此，本文不能把 dual-gated bilayer WSe2、layer-dependent exciton-polaron、低场 SHG 增强或普通 SHG 开关写成自己的创新。这些是基线工作已经建立的实验平台和强度响应。本文的创新边界是：在同一类双栅 exciton-polaron 平台上，不再只研究 SHG intensity，而是研究复数 SHG 场 E2w 本身是否可以通过门压主动推过相干零点，并由此产生相位跳变和绕转。",
    )

    doc.add_heading("2. 文献缺口：从强度控制到复场控制", level=1)
    add_paragraph(
        doc,
        "电控 SHG 强度并不是空白问题。Seyler 等人在 monolayer WSe2 transistor 中已经展示了电掺杂对二阶非线性响应的调控，低温下 A-exciton 附近 SHG 强度可调超过一个数量级，室温下也接近四倍[2]。这说明“电控 SHG 强度”本身不足以构成本文的核心创新。",
    )
    add_paragraph(
        doc,
        "另一方面，SHG 场具有复相位并可发生层间相干干涉也已有先例。Kim 等人的 Second-Harmonic Young Interference 工作用 spectral phase interferometry 证明了 atom-thin heterocrystals 中两层发出的二次谐波场会相干干涉，并可解析材料相关的 SH phase delay[3]。这说明复数二阶非线性极化率和 SHG 相位并不是形式上的数学量，而是可以实验访问的物理量。",
    )
    add_paragraph(
        doc,
        "真正的缺口是：已有工作要么关注电控强度，要么关注静态材料组合或扭角导致的 SHG 相位干涉。仍然缺少一种以双栅为主动控制参数、以 exciton-polaron 共振为相位旋转源、以复 SHG 场零点和相位绕转为核心判据的机制方案。本文正是在这个缺口上提出并数值认证一个单一创新点。",
    )

    doc.add_heading("3. 核心物理命题", level=1)
    add_paragraph(
        doc,
        "本文的核心命题可以写成一句话：双栅调控 exciton-polaron 共振，实现双层 WSe2 标量复 SHG 场的电控零点与相位绕转。这里的“标量复 SHG 场”指固定入射和探测偏振通道下的 E2w 复振幅，而不是完整 Jones vector 偏振态。这样做的好处是主线集中、实验判据清晰，也避免把本文扩展成 polarization singularity 的主张。",
    )
    add_paragraph(
        doc,
        "在 2H bilayer WSe2 中，上下两层二阶响应近似反号。若门压使上层和下层的复数非线性极化率分别为 chi_top 和 chi_bottom，则最小模型中的发射场可以写成 E2w = chi_top - chi_bottom。普通强度实验只看 |E2w|^2，而相位分辨实验可以恢复 arg(E2w) 以及 Re(E2w)、Im(E2w)。当 Re(E2w) 和 Im(E2w) 在二维门压参数空间中同时过零时，就形成一个标量复场零点。",
    )
    add_paragraph(
        doc,
        "双栅参数空间提供两个独立控制量，通常可重写为 carrier density n 和 displacement field D。复场零点需要满足两个方程 Re(E2w)=0 与 Im(E2w)=0，因此在二维参数空间中出现孤立零点是自然的。围绕该零点走闭合 gate loop 时，arg(E2w) 的总变化可以接近 2pi，对应非零 winding number。这正是本文把强度调制升级为复场拓扑响应的关键。",
    )

    doc.add_heading("4. 可复现 Python 模型", level=1)
    add_paragraph(
        doc,
        "Python 模型采用定性到半定量之间的 literature-constrained phenomenological 形式。它不拟合 Cha 2024 原始实验数据，也不试图替代第一性原理或电磁场仿真。模型的目标是数值认证机制链条是否自洽：gate 改变层分辨载流子分布，载流子分布改变 neutral-exciton 与 exciton-polaron 的复 Lorentz 响应，层分辨复响应相干叠加后产生接近零点的标量复 SHG 场。",
    )
    add_paragraph(
        doc,
        "模型分为五个模块。第一，gate coordinate 模块将 Vt、Vb 映射为归一化 n 和 D。第二，susceptibility 模块为每一层构造 neutral-exciton 与 exciton-polaron 两个复 Lorentz 振子，并允许门压导致振子强度、线宽和共振位置变化。第三，SHG field 模块实现 E2w = chi_top - chi_bottom 的 2H 双层相干场求和。第四，topology 模块寻找近零点并计算闭合路径相位绕数。第五，heterodyne 模块模拟样品 SHG 与参考 SHG 的干涉条纹。",
    )
    add_paragraph(
        doc,
        "为避免漂亮但不可测的数值结果，模型加入轻量实验可实现约束：门压和位移场保持在归一化合理范围内，路径上除接近零点的区域外保持可检测相对强度，线宽和振子强度的扫描避免极端值。本文不做完整 device optimizer，而是用这些 filter 排除明显不合理的参数选择。",
    )

    doc.add_heading("5. 数值认证结果", level=1)
    add_figure(
        doc,
        "figure_1_setup_coordinates.png",
        "图 1 给出创新机制起点和模型坐标。它不是 Cha 2024 原图复现，而是说明本文如何从双栅、层选择性 exciton-polaron 和 2H 相干场求和出发，转向 Re(E)、Im(E)、|E|^2 与 arg(E) 等复场观测量。",
    )
    add_figure(
        doc,
        "figure_2_complex_field_maps.png",
        "图 2 展示标量复 SHG 场的多面板图：归一化强度、相位、实部和虚部。近零点附近强度出现暗点，同时相位发生快速变化；Re(E) 与 Im(E) 的过零结构共同定义了复场零点。",
    )
    add_figure(
        doc,
        "figure_3_gate_loop_winding.png",
        "图 3 选择围绕近零点的闭合 gate loop，展示门压路径、复平面轨迹、路径上的相对强度和展开相位。复平面轨迹围绕原点时，相位沿路径累计变化接近 2pi，形成数值认证的 phase winding。",
    )
    add_figure(
        doc,
        "figure_4_robustness.png",
        "图 4 用轻量参数扫描检查现象是否依赖单点调参。通过改变 exciton-polaron 线宽和振子强度，可以观察近零点深度和 loop winding 的稳定区域。该图的目的不是完整优化，而是说明机制在一片合理参数区间内存在。",
    )
    add_figure(
        doc,
        "figure_5_heterodyne.png",
        "图 5 将复场相位转换为 heterodyne SHG 可测量量。样品 SHG 与参考 SHG 干涉后，gate loop 上的相位变化表现为干涉条纹的整体平移和相位读出曲线的连续绕转。这为实验验证提供直接判据。",
    )

    doc.add_heading("6. 实验验证方案与对照", level=1)
    add_paragraph(
        doc,
        "实验上，主样品仍可采用 hBN-encapsulated 2H bilayer WSe2 with graphite top and bottom gates。普通 SHG gate map 只能验证强度随门压变化，而本文真正需要的是 phase-resolved SHG 或 heterodyne SHG。样品 SHG 场与 quartz、BBO 或参考单层 WSe2 产生的参考 SHG 场干涉，通过扫描 reference phase 或 delay 恢复样品 E2w 的相位。",
    )
    add_paragraph(
        doc,
        "关键判据有四个。第一，沿特定 gate path 接近零点时 |E2w|^2 出现强度极小。第二，相同位置附近 Re(E2w) 与 Im(E2w) 同时接近零。第三，穿过零点附近时相位出现接近 pi 的跳变。第四，围绕零点闭合扫描时相位累计变化接近 2pi。只有同时满足这些相位敏感判据，才能区分真正的相干零点和普通强度淬灭。",
    )
    add_paragraph(
        doc,
        "必要对照包括远离 exciton-polaron 共振的 pump detuning、monolayer WSe2 参考、以及 D 反向时上下层角色交换导致零点位置镜像移动。远离共振后相位绕转应减弱或消失；单层样品缺少双层相干抵消通道；门压极性反转应改变层选择性极化子的空间分布。这些对照可以排除吸收、加热或普通 oscillator strength 变化导致的假零点解释。",
    )

    doc.add_heading("7. 模型边界与可验证性", level=1)
    add_bullets(
        doc,
        [
            "本文不拟合 Cha 2024 原始实验数据；Python 结果是创新机制的数值认证，而不是对实验图的复现。",
            "本文不声称已经实验观测到 SHG zero 或 phase winding；实验部分是 proposed validation path。",
            "本文不做 DFT、COMSOL 或 FDTD。核心问题是复 SHG 场的半经验相干机制，不是第一性原理材料参数提取或电磁场分布仿真。",
            "本文不把 polarization singularity / Poincare sphere 作为 WSe2 主 claim。相关内容只作为 twisted heterobilayer 或偏振分辨扩展方向。",
        ],
    )
    add_paragraph(
        doc,
        "这些边界并不削弱本文的核心结论。相反，它们让本文的主张更明确：在已建立的双栅 exciton-polaron SHG 平台上，复数二次谐波场的零点和相位绕转是一个可以被数值认证、并可由 heterodyne SHG 直接检验的具体创新点。",
    )

    doc.add_heading("8. 结论", level=1)
    add_paragraph(
        doc,
        "本文将主论文中的 layer-dependent exciton-polaron-enhanced SHG 从强度调制问题推进为标量复 SHG 场控制问题。通过 literature-constrained Python 模型，我们数值认证了双栅调控可以把 E2w 推近相干零点，并在闭合 gate loop 上产生相位绕转。这个结果给出了一个明确的创新路线：不是继续追求更大的 SHG enhancement，而是用双栅把 exciton-polaron 态变成可编程的非线性相位源。",
    )

    doc.add_heading("参考文献", level=1)
    refs = [
        "S. Cha, T. Ouyang, T. Taniguchi, K. Watanabe, N. M. Gabor, and C. H. Lui, Enhancing Resonant Second-Harmonic Generation in Bilayer WSe2 by Layer-Dependent Exciton-Polaron Effect, Nano Letters 24, 14847-14853 (2024). DOI: 10.1021/acs.nanolett.4c04544.",
        "K. L. Seyler et al., Electrical Control of Second-Harmonic Generation in a WSe2 Monolayer Transistor, Nature Nanotechnology 10, 407-411 (2015). DOI: 10.1038/nnano.2015.73.",
        "W. Kim, J. Y. Ahn, J. Oh, J. H. Shim, and S. Ryu, Second-Harmonic Young Interference in Atom-Thin Heterocrystals, Nano Letters 20, 8825 (2020). DOI: 10.1021/acs.nanolett.0c03763.",
        "Y. Wang et al., Destructive interference of second harmonic generation in AA stacked MoTe2/WSe2, arXiv:2605.21231 (2026).",
        "R. Zu et al., Analytical and numerical modeling of optical second harmonic generation in anisotropic crystals using #SHAARP package, npj Computational Materials 8, 246 (2022). DOI: 10.1038/s41524-022-00930-4.",
    ]
    for idx, ref in enumerate(refs, 1):
        doc.add_paragraph(f"[{idx}] {ref}")

    doc.save(PRIMARY_DOC)
    shutil.copyfile(PRIMARY_DOC, ASCII_DOC)


if __name__ == "__main__":
    build_document()

